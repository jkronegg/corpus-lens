#!/usr/bin/env python3
"""Exporte la consommation Romande Energie en totaux mensuels normalisés vers Excel.

Normalisation appliquée pour chaque mois:
- total_normalise = (total_mesure / jours_de_donnees) * jours_du_mois
"""

from __future__ import annotations

import argparse
import calendar
import csv
import datetime as dt
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

DATE_FORMAT = "%d.%m.%Y %H:%M:%S"
DEFAULT_DELIMITER = ";"


@dataclass
class MonthlyStats:
    total: float = 0.0
    days_with_data: set[dt.date] | None = None

    def __post_init__(self) -> None:
        if self.days_with_data is None:
            self.days_with_data = set()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Agrège un CSV Romande Energie quart-horaire/horaire en consommation mensuelle "
            "et exporte un fichier Excel avec normalisation par nombre de jours."
        )
    )
    parser.add_argument("input_csv", help="Chemin du CSV d'entree")
    parser.add_argument(
        "--output-xlsx",
        help="Chemin du fichier Excel de sortie (defaut: <input>_monthly_normalized.xlsx)",
    )
    parser.add_argument(
        "--delimiter",
        default=DEFAULT_DELIMITER,
        help=f"Separateur CSV (defaut: {DEFAULT_DELIMITER!r})",
    )
    parser.add_argument("--date-column", default="Date", help="Nom de la colonne date")
    parser.add_argument(
        "--consumption-column",
        default=None,
        help="Nom de la colonne consommation (defaut: detection automatique)",
    )
    return parser.parse_args()


def detect_consumption_column(fieldnames: list[str], requested: str | None) -> str:
    if requested:
        if requested not in fieldnames:
            raise SystemExit(f"Colonne consommation introuvable: {requested!r}")
        return requested

    for name in fieldnames:
        if "consommation" in name.lower():
            return name

    if len(fieldnames) >= 2:
        return fieldnames[1]

    raise SystemExit("Impossible de detecter la colonne consommation.")


def parse_float(value: str) -> float:
    text = (value or "").strip().replace(" ", "")
    if not text:
        return 0.0
    text = text.replace(",", ".")
    return float(text)


def aggregate_monthly(
    input_csv: Path,
    delimiter: str,
    date_column: str,
    consumption_column: str | None,
) -> list[tuple[int, int, MonthlyStats]]:
    monthly: dict[tuple[int, int], MonthlyStats] = {}

    with input_csv.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter=delimiter)
        if not reader.fieldnames:
            raise SystemExit(f"CSV vide ou en-tete invalide: {input_csv}")
        if date_column not in reader.fieldnames:
            raise SystemExit(f"Colonne date introuvable: {date_column!r}")

        conso_col = detect_consumption_column(reader.fieldnames, consumption_column)

        for row in reader:
            date_raw = (row.get(date_column) or "").strip()
            conso_raw = (row.get(conso_col) or "").strip()
            if not date_raw:
                continue

            timestamp = dt.datetime.strptime(date_raw, DATE_FORMAT)
            value = parse_float(conso_raw)

            key = (timestamp.year, timestamp.month)
            stats = monthly.setdefault(key, MonthlyStats())
            stats.total += value
            stats.days_with_data.add(timestamp.date())

    return [(year, month, monthly[(year, month)]) for (year, month) in sorted(monthly)]


def default_output_path(input_csv: Path) -> Path:
    return input_csv.with_name(f"{input_csv.stem}_monthly_normalized.xlsx")


def write_excel(rows: list[tuple[int, int, MonthlyStats]], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Conso mensuelle"

    headers = [
        "Mois",
        "Total mesure (kWh)",
        "Jours de donnees",
        "Jours du mois",
        "Total normalise (kWh)",
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for year, month, stats in rows:
        days_with_data = len(stats.days_with_data)
        days_in_month = calendar.monthrange(year, month)[1]
        normalized = 0.0
        if days_with_data > 0:
            normalized = (stats.total / days_with_data) * days_in_month

        ws.append(
            [
                f"{year:04d}-{month:02d}",
                round(stats.total, 3),
                days_with_data,
                days_in_month,
                round(normalized, 3),
            ]
        )

    for col_idx in range(1, 6):
        width = 20 if col_idx in (2, 5) else 16
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "0.000"
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = "0.000"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> None:
    args = parse_args()
    input_csv = Path(args.input_csv).resolve()

    if not input_csv.exists():
        raise SystemExit(f"CSV introuvable: {input_csv}")

    output_xlsx = Path(args.output_xlsx).resolve() if args.output_xlsx else default_output_path(input_csv)
    rows = aggregate_monthly(
        input_csv=input_csv,
        delimiter=args.delimiter,
        date_column=args.date_column,
        consumption_column=args.consumption_column,
    )
    write_excel(rows, output_xlsx)

    print(f"Excel genere: {output_xlsx}")
    print(f"Mois exportes: {len(rows)}")


if __name__ == "__main__":
    main()


