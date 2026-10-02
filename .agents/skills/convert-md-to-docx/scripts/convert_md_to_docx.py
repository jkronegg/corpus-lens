#!/usr/bin/env python3
"""Convertir un fichier Markdown en document Word (.docx)."""

from __future__ import annotations

import argparse
import os
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Cm


# Décale toutes les listes d'environ 1 cm vers la droite.
LIST_BASE_INDENT_CM = 1.0
DEFAULT_AUTHOR = os.getenv('CORPUS_LENS_AUTHOR', 'Julien')


def normalize_metadata_text(value: str | None, max_length: int = 255) -> str:
    """Nettoyer une valeur de métadonnée en évitant les chaînes trop longues."""
    if not value:
        return ''
    cleaned = re.sub(r'\s+', ' ', value).strip()
    if len(cleaned) > max_length:
        cleaned = cleaned[: max_length - 3].rstrip() + '...'
    return cleaned


def parse_front_matter(content: str) -> dict[str, object]:
    """Extraire les métadonnées YAML en tête de document si présentes."""
    metadata: dict[str, object] = {}
    if not content.startswith('---'):
        return metadata

    lines = content.splitlines()
    if len(lines) < 3:
        return metadata

    end_index = None
    for idx in range(1, len(lines)):
        if lines[idx].strip() == '---':
            end_index = idx
            break
    if end_index is None:
        return metadata

    for raw_line in lines[1:end_index]:
        if ':' not in raw_line:
            continue
        key, value = raw_line.split(':', 1)
        key = key.strip()
        value = value.strip()
        if not key:
            continue
        if value in {'', "''", '""'}:
            metadata[key.lower()] = ''
        else:
            metadata[key.lower()] = value.strip("\"'")
    return metadata


def strip_front_matter(content: str) -> str:
    """Retourner le corps Markdown sans le front matter YAML initial."""
    if not content.startswith('---'):
        return content

    lines = content.splitlines()
    if len(lines) < 3:
        return content

    for idx in range(1, len(lines)):
        if lines[idx].strip() == '---':
            return '\n'.join(lines[idx + 1 :])
    return content


def is_blockquote_line(line: str) -> bool:
    """Retourner True si la ligne Markdown appartient a un bloc citation."""
    return line.startswith('>')


def get_blockquote_content(line: str) -> str:
    """Extraire le contenu texte d'une ligne de citation (`>` optionnellement suivi d'un espace)."""
    content = line[1:]
    if content.startswith(' '):
        return content[1:]
    return content


def extract_markdown_title(lines: list[str], front_matter: dict[str, object] | None = None) -> str:
    """Extraire le titre principal du Markdown, ou une valeur de secours."""
    if front_matter:
        for key in ('title', 'document_title', 'nom'):
            value = front_matter.get(key)
            if isinstance(value, str) and value.strip():
                return value.strip()

    for line in lines:
        stripped = line.strip()
        if stripped.startswith('# '):
            return stripped[2:].strip()
        if stripped.startswith('## '):
            return stripped[3:].strip()
    for line in lines:
        stripped = line.strip()
        if stripped and not stripped.startswith('#') and not stripped.startswith('>'):
            return stripped
    return 'Document'


def extract_document_summary(lines: list[str], title: str, front_matter: dict[str, object] | None = None) -> str:
    """Créer un résumé à partir du texte du document."""
    if front_matter:
        for key in ('summary', 'description', 'abstract', 'resume'):
            value = front_matter.get(key)
            if isinstance(value, str) and value.strip():
                return normalize_metadata_text(value, 255)

    candidates: list[str] = []
    in_keywords_section = False

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.lower().startswith('# mots clés') or stripped.lower().startswith('## mots clés'):
            in_keywords_section = True
            continue
        if in_keywords_section:
            if stripped.startswith('- ') or stripped.startswith('* '):
                continue
            if stripped.startswith('#'):
                in_keywords_section = False
        if stripped.startswith('#'):
            continue
        if stripped.startswith('>'):
            continue
        if stripped.startswith('- ') or stripped.startswith('* '):
            continue
        candidates.append(stripped)
        if len(candidates) >= 3:
            break

    if candidates:
        summary = ' '.join(candidates)
        summary = re.sub(r'\s+', ' ', summary)
        return normalize_metadata_text(summary, 255)

    if title and title != 'Document':
        return normalize_metadata_text(f"Document consacré à {title}.", 255)
    return 'Document préparé pour consultation et analyse.'


def extract_document_keywords(lines: list[str], title: str, front_matter: dict[str, object] | None = None) -> list[str]:
    """Extraire les mots-clés depuis le front matter ou la section Markdown dédiée."""
    if front_matter:
        for key in ('keywords', 'tags', 'mots_cles', 'mots-clés'):
            value = front_matter.get(key)
            if isinstance(value, str):
                parsed = [item.strip().strip(',.;:!') for item in value.split(',') if item.strip()]
                if parsed:
                    return parsed[:10]
            elif isinstance(value, list):
                parsed = [str(item).strip() for item in value if str(item).strip()]
                if parsed:
                    return parsed[:10]

    keywords: list[str] = []
    in_keywords_section = False

    for line in lines:
        stripped = line.strip()
        if not stripped:
            if in_keywords_section:
                break
            continue
        if stripped.lower().startswith('# mots clés') or stripped.lower().startswith('## mots clés'):
            in_keywords_section = True
            continue
        if in_keywords_section:
            if stripped.startswith('- ') or stripped.startswith('* '):
                item = stripped[2:].strip()
                if item:
                    keywords.append(item)
                continue
            if stripped.startswith('#'):
                break
        if stripped.startswith('- ') or stripped.startswith('* '):
            item = stripped[2:].strip()
            if item:
                keywords.append(item)

    if keywords:
        unique_keywords: list[str] = []
        seen: set[str] = set()
        for keyword in keywords:
            normalized = keyword.strip().strip(',.;:!')
            if normalized and normalized.lower() not in seen:
                unique_keywords.append(normalized)
                seen.add(normalized.lower())
        return unique_keywords[:10]

    fallback = re.split(r'\s+|[-–—]', title.strip()) if title else []
    clean_fallback = [token.strip('.,;:!?”\'()[]{}') for token in fallback if token and len(token) > 2]
    if clean_fallback:
        return clean_fallback[:8]
    return ['document', 'markdown', 'analyse']


def extract_document_author(front_matter: dict[str, object] | None = None) -> str:
    """Extraire le nom de l'auteur depuis le front matter ou la configuration de l'environnement."""
    if front_matter:
        for key in ('author', 'auteur', 'creator', 'created_by'):
            value = front_matter.get(key)
            if isinstance(value, str) and value.strip():
                return value.strip()
    return DEFAULT_AUTHOR


def format_quote_paragraph(paragraph) -> None:
    """Appliquer la mise en forme de citation Word sur un paragraphe."""
    paragraph.paragraph_format.left_indent = 720000
    paragraph.paragraph_format.space_before = 120000
    paragraph.paragraph_format.space_after = 120000


def add_hyperlink(paragraph, text: str, url: str, force_italic: bool = False) -> None:
    """Ajouter un hyperlien cliquable à un paragraphe Word."""
    part = paragraph.part
    rel_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), rel_id)

    run = OxmlElement('w:r')
    run_props = OxmlElement('w:rPr')

    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0000FF')
    run_props.append(color)

    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    run_props.append(underline)

    if force_italic:
        italic = OxmlElement('w:i')
        run_props.append(italic)

    run.append(run_props)

    text_node = OxmlElement('w:t')
    text_node.text = text
    run.append(text_node)

    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_formatted_text(paragraph, text: str, force_italic: bool = False) -> None:
    """Ajouter du texte formaté (gras, italique, liens, code) à un paragraphe."""
    pattern = r'(\*\*.*?\*\*|\*.*?\*|\[.*?\]\(.*?\)|`.*?`)' 
    parts = re.split(pattern, text)

    for part in parts:
        if not part:
            continue

        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.italic = force_italic
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.italic = force_italic
        elif part.startswith('[') and '](' in part:
            match = re.match(r'\[(.*?)\]\((.*?)\)', part)
            if match:
                text_link = match.group(1)
                url = match.group(2)
                add_hyperlink(paragraph, text_link, url, force_italic=force_italic)
            else:
                run = paragraph.add_run(part)
                run.italic = force_italic
        else:
            run = paragraph.add_run(part)
            run.italic = force_italic


def parse_markdown_to_docx(md_file: str, docx_file: str) -> None:
    """Convertir un Markdown en document Word."""
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()

    front_matter = parse_front_matter(content)
    markdown_body = strip_front_matter(content)
    lines = markdown_body.split('\n')
    title = extract_markdown_title(lines, front_matter)
    summary = extract_document_summary(lines, title, front_matter)
    keywords = extract_document_keywords(lines, title, front_matter)
    author = extract_document_author(front_matter)

    doc = Document()
    doc.core_properties.author = author
    doc.core_properties.title = normalize_metadata_text(title, 255)
    doc.core_properties.subject = summary
    doc.core_properties.keywords = ', '.join(keywords)

    paragraph_buffer: list[str] = []

    def flush_paragraph_buffer() -> None:
        """Écrire le paragraphe texte en attente (lignes Markdown fusionnées)."""
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        p = doc.add_paragraph()
        add_formatted_text(p, ' '.join(paragraph_buffer))
        paragraph_buffer = []

    i = 0
    while i < len(lines):
        line = lines[i]

        if line.startswith('# '):
            flush_paragraph_buffer()
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith('## '):
            flush_paragraph_buffer()
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith('### '):
            flush_paragraph_buffer()
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith('#### '):
            flush_paragraph_buffer()
            doc.add_heading(line[5:].strip(), level=4)
        elif is_blockquote_line(line):
            flush_paragraph_buffer()
            quote_paragraph_lines: list[str] = []

            def flush_quote_paragraph() -> None:
                """Ecrire un paragraphe de citation en joignant ses lignes."""
                nonlocal quote_paragraph_lines
                if not quote_paragraph_lines:
                    return
                p = doc.add_paragraph()
                format_quote_paragraph(p)
                add_formatted_text(p, ' '.join(quote_paragraph_lines), force_italic=True)
                quote_paragraph_lines = []

            while i < len(lines) and is_blockquote_line(lines[i]):
                quote_line = get_blockquote_content(lines[i]).strip()
                if quote_line:
                    quote_paragraph_lines.append(quote_line)
                else:
                    flush_quote_paragraph()
                i += 1
            i -= 1
            flush_quote_paragraph()
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            flush_paragraph_buffer()
            indent_level = (len(line) - len(line.lstrip())) // 2
            text = line.strip()[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Cm(LIST_BASE_INDENT_CM + indent_level * 2.0)
            add_formatted_text(p, text)
        elif re.match(r'^\d+\)', line.strip()):
            flush_paragraph_buffer()
            indent_level = (len(line) - len(line.lstrip())) // 2
            text = re.sub(r'^\d+\)\s*', '', line.strip())
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.left_indent = Cm(LIST_BASE_INDENT_CM + indent_level * 2.0)
            add_formatted_text(p, text)
        elif not line.strip():
            flush_paragraph_buffer()
        else:
            paragraph_buffer.append(line.strip())

        i += 1

    flush_paragraph_buffer()

    doc.save(docx_file)
    print(f"Document créé : {docx_file}")


def infer_output_path(md_path: Path) -> Path:
    return md_path.with_suffix('.docx')


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Convertir un fichier Markdown en document Word (.docx)."
    )
    parser.add_argument(
        '--input',
        required=True,
        help='Chemin du fichier Markdown source.',
    )
    parser.add_argument(
        '--output',
        help='Chemin du fichier Word de sortie (.docx). Si omis, le même nom est utilisé.',
    )
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    md_path = Path(args.input).expanduser().resolve()
    if not md_path.exists():
        parser.error(f"Fichier Markdown introuvable: {md_path}")

    docx_path = Path(args.output).expanduser().resolve() if args.output else infer_output_path(md_path)
    parse_markdown_to_docx(str(md_path), str(docx_path))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())

